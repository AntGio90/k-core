import os
from datetime import datetime
import pytz
from docx import Document as DocxDocument
from cat.mad_hatter.decorators import hook, tool
from cat.convo.messages import CatMessage
import logging

# Set up logging
logging.basicConfig(level=logging.INFO)

@hook(priority=10)
def before_cat_sends_message(message: CatMessage, cat) -> CatMessage:
    """This hook checks if the user asked to save the response and triggers the save functionality."""
    
    # Get the user's last message from working memory
    if hasattr(cat, "working_memory") and hasattr(cat.working_memory, "user_message_json"):
        user_message = cat.working_memory.user_message_json.text.lower()
        
        # Check if the user asked to save the response
        keywords = ["write", "save", "store", "export", "document", "record", "keep"]
        if any(keyword in user_message for keyword in keywords):
            return store_output(message, cat)
    
    return message

def store_output(message: CatMessage, cat) -> CatMessage:
    """Saves the cat's response to a Word document."""
    
    timestamp = datetime.now(pytz.timezone('Europe/Rome')).strftime("%Y%m%d_%H%M%S")
    filename = f"cheshire_cat_response_{timestamp}.docx"
    
    # Save in the plugin's directory
    plugin_dir = os.path.dirname(__file__)
    
    # Create directory if it doesn't exist
    if not os.path.exists(plugin_dir):
        try:
            os.makedirs(plugin_dir)
            logging.info(f"Created directory: {plugin_dir}")
        except Exception as e:
            logging.error(f"Failed to create directory: {e}")
            return message
    
    filepath = os.path.join(plugin_dir, filename)

    # Try to create and save the document
    try:
        logging.info(f"Attempting to save document to: {filepath}")
        logging.info(f"Message content: {message.text}")
        
        docx_doc = DocxDocument()
        docx_doc.add_heading("Cheshire Cat AI Response", 0)
        docx_doc.add_paragraph(message.text)    
        docx_doc.save(filepath)
        
        logging.info(f"Document saved successfully to: {filepath}")
        
        # Append notification to the message
        notification = f"\n\n_Response saved at: {filepath}_"
        message.text = message.text + notification
        
    except Exception as e:
        logging.error(f"Error saving document: {e}")
        message.text = message.text + f"\n\n_Failed to save response: {str(e)}_"
        
    return message